from typing import List
from docx import Document
from langchain.vectorstores import FAISS
from langchain_community.embeddings import SentenceTransformerEmbeddings

# Initialize a pre-trained model for embeddings
embedding_model = SentenceTransformerEmbeddings(model_name='quentinz/bge-large-zh-v1.5:latest',host='http://localhost:11434')

def extract_text_and_images(docx_path: str):
    """Extract text and image placeholders from a Word document."""
    doc = Document(docx_path)
    texts = []
    for para in doc.paragraphs:
        if para.text.strip():
            texts.append(para.text.strip())
    
    # For simplicity, we return image placeholders (implement image embeddings as needed)
    images = ["Image Placeholder"] * len(doc.inline_shapes)  # Placeholder for image handling
    
    return texts, images

def store_to_faiss(data: List[str], index_path: str):
    """Store data in a Faiss index using LangChain."""
    # Compute embeddings using LangChain
    vectorstore = FAISS.from_texts(data, embedding_model)
    # Save the index to disk
    vectorstore.save_local(index_path)

def load_from_faiss(index_path: str):
    """Load a Faiss index using LangChain."""
    return FAISS.load_local(index_path, embedding_model)

def search_in_faiss(vectorstore, query: str, top_k: int = 5):
    """Search for the most similar embeddings in the Faiss index."""
    results = vectorstore.similarity_search(query, k=top_k)
    return results

if __name__ == "__main__":
    # Example usage
    docx_path = "RMAN.docx"
    index_path = "faiss_index"

    # Step 1: Extract text and images from the Word document
    texts, images = extract_text_and_images(docx_path)
    combined_data = texts + images

    # Step 2: Store data to Faiss
    store_to_faiss(combined_data, index_path)

    # Step 3: Load data from Faiss
    vectorstore = load_from_faiss(index_path)

    # Step 4: Search for a query
    query = "Sample text from the document"
    results = search_in_faiss(vectorstore, query)

    # Print results
    for result in results:
        print("Document:", result.page_content)
